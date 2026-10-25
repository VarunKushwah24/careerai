import os
from pathlib import Path
from docx import Document

def txt_to_docx(txt_path: Path, docx_path: Path):
    if not txt_path.exists():
        print(f"Not found: {txt_path}")
        return
    with open(txt_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if stripped == stripped.upper() and len(stripped) < 40 and not '@' in stripped:
            doc.add_heading(stripped, level=1)
        elif stripped.startswith('-') or stripped.startswith('*'):
            doc.add_paragraph(stripped.lstrip('-* ').strip(), style='List Bullet')
        else:
            doc.add_paragraph(stripped)
            
    doc.save(str(docx_path))
    print(f"Generated: {docx_path}")

if __name__ == '__main__':
    base = Path(__file__).parent
    txt_to_docx(base / 'Software_Engineer_Resume.txt', base / 'Software_Engineer_Resume.docx')
    txt_to_docx(base / 'Data_Scientist_Resume.txt', base / 'Data_Scientist_Resume.docx')
