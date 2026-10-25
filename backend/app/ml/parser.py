import re
from pathlib import Path
from typing import Optional
from pypdf import PdfReader
from docx import Document

_ocr_engine = None

def get_ocr_engine():
    """Lazily load RapidOCR engine with cached model weights in memory."""
    global _ocr_engine
    if _ocr_engine is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
            _ocr_engine = RapidOCR()
        except Exception as e:
            print(f"Warning: Could not initialize RapidOCR engine: {e}")
            _ocr_engine = None
    return _ocr_engine

def extract_text_with_ocr(image_source) -> str:
    """
    Extract text from an image source (PIL Image, numpy array, or file path) using RapidOCR.
    Sorts detected text blocks top-to-bottom, left-to-right to preserve document hierarchy.
    """
    engine = get_ocr_engine()
    if engine is None:
        return ""
    try:
        from PIL import Image
        import numpy as np

        if isinstance(image_source, Image.Image):
            img_arr = np.array(image_source.convert("RGB"))
        elif isinstance(image_source, (str, Path)):
            with Image.open(str(image_source)) as img:
                img_arr = np.array(img.convert("RGB"))
        elif isinstance(image_source, bytes):
            import io
            with Image.open(io.BytesIO(image_source)) as img:
                img_arr = np.array(img.convert("RGB"))
        elif isinstance(image_source, np.ndarray):
            img_arr = image_source
        else:
            return ""

        result, _ = engine(img_arr)
        if not result:
            return ""

        # RapidOCR returns list of: [box, text, score]
        # Sort text blocks top-to-bottom (min y) then left-to-right (min x)
        sorted_blocks = sorted(
            result,
            key=lambda item: (min(p[1] for p in item[0]), min(p[0] for p in item[0]))
        )
        lines = [item[1].strip() for item in sorted_blocks if item and len(item) > 1 and item[1].strip()]
        return "\n".join(lines)
    except Exception as e:
        print(f"OCR execution warning: {e}")
        return ""

def extract_text_from_image(file_path: str) -> str:
    """Extract text from standalone image files (PNG, JPG, JPEG, WEBP, BMP, TIFF) using AI OCR."""
    return extract_text_with_ocr(file_path)

def extract_text_from_pdf(file_path: str) -> str:
    """
    Extract text from PDF documents.
    Fast path: Extracts digital text stream via pypdf.
    Fallback (Scanned PDF): If digital text < 50 chars, renders pages via pypdfium2 and runs AI OCR.
    """
    text_chunks = []
    try:
        with open(file_path, "rb") as f:
            reader = PdfReader(f)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text and page_text.strip():
                    text_chunks.append(page_text.strip())
    except Exception as e:
        print(f"Warning reading PDF digital text: {str(e)}")

    extracted = "\n".join(text_chunks)

    # Fast path: If digital text stream has sufficient characters, return immediately
    if extracted and len(extracted.strip()) >= 50:
        return extracted

    # Scanned PDF Fallback: Render each page to an image and run OCR
    ocr_chunks = []
    try:
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(file_path)
        try:
            for page_idx in range(len(pdf)):
                page = pdf[page_idx]
                # Render at 2.0x scale (144 DPI) for crisp character recognition
                pil_image = page.render(scale=2.0).to_pil()
                page_text = extract_text_with_ocr(pil_image)
                if page_text and page_text.strip():
                    ocr_chunks.append(page_text.strip())
        finally:
            pdf.close()
    except Exception as pdfium_err:
        print(f"pypdfium2 OCR fallback error: {pdfium_err}")

    # Secondary fallback: Extract embedded images from pypdf if rendering failed
    if not ocr_chunks:
        try:
            with open(file_path, "rb") as f:
                reader = PdfReader(f)
                for page in reader.pages:
                    for img in getattr(page, "images", []):
                        img_text = extract_text_with_ocr(img.data)
                        if img_text and img_text.strip():
                            ocr_chunks.append(img_text.strip())
        except Exception as img_err:
            print(f"pypdf embedded images OCR fallback error: {img_err}")

    ocr_extracted = "\n".join(ocr_chunks)
    if ocr_extracted and len(ocr_extracted.strip()) > len(extracted.strip()):
        return ocr_extracted

    return extracted

def extract_text_from_docx(file_path: str) -> str:
    text_chunks = []
    try:
        doc = Document(file_path)
        for p in doc.paragraphs:
            if p.text.strip():
                text_chunks.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_chunks.append(" | ".join(row_text))
    except Exception as e:
        raise ValueError(f"Error reading DOCX file: {str(e)}")
    return "\n".join(text_chunks)

def clean_text(raw_text: str) -> str:
    if not raw_text:
        return ""
    # Normalize bullet points and weird dashes
    text = re.sub(r"[\u2022\u2023\u25E6\u2043\u2219\u00B7\u25AA\u25CF]", "\n* ", raw_text)
    text = re.sub(r"[\u2013\u2014]", "-", text)
    # Remove null bytes or non-printable chars
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]", " ", text)
    # Normalize whitespace while preserving linebreaks
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    # Remove consecutive empty lines
    cleaned = "\n".join([line for i, line in enumerate(lines) if line or (i > 0 and lines[i-1])])
    return cleaned.strip()

def parse_resume_file(file_path: str) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        raw = extract_text_from_pdf(file_path)
    elif suffix in [".docx", ".doc"]:
        raw = extract_text_from_docx(file_path)
    elif suffix in [".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"]:
        raw = extract_text_from_image(file_path)
    else:
        raise ValueError(
            f"Unsupported file format: {suffix}. Supported formats are PDF, DOCX, and Images (PNG, JPG, WEBP, BMP, TIFF)."
        )
    
    cleaned = clean_text(raw)
    if not cleaned or len(cleaned.strip()) < 30:
        raise ValueError(
            "Could not extract sufficient text from resume. Please ensure the document or image is clear and contains readable text."
        )
    return cleaned
